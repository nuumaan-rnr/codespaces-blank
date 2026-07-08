"""RSTAB-matching config options: EN1993 imperfection, connector override,
nonlinear axial-dependent base."""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from rack15512.builder import RackConfig, build_rack, LevelSpec
from rack15512.engine.opensees import OpenSeesEngine
from rack15512.model import Imperfection


def test_en1993_imperfection_is_flat_phi_s():
    imp = Imperfection(n_cols=4, phi_s=1/300, standard="EN1993", phi_min=1/500)
    assert abs(imp.value() - 1/300) < 1e-9          # flat, no 2x / sqrt factor
    imp15 = Imperfection(n_cols=4, phi_s=1/300, standard="EN15512", phi_l=0)
    assert imp15.value() > imp.value()              # EN 15512 is larger (amplified)


def test_en1993_alpha_hm_reduction_matches_rstab():
    # RSTAB "Calculate value of inclination" (EN 1993-1-1 5.3.2(3)):
    # Phi = Phi0 * alpha_h * alpha_m, with h = 7965 mm, m = 4 columns/row.
    m = build_rack(RackConfig(module="single", n_bays=3,
                              levels=[LevelSpec(gap=1800.0)], frame_height=7965.0,
                              imperfection_standard="EN1993",
                              imperfection_alpha_hm=True,
                              phi_s=1/200, phi_s_cross=1/300))
    imp = m.imperfection
    assert imp.n_cols == 4 and imp.alpha_hm and imp.height == 7965.0
    assert abs(imp._alpha_hm() - 0.5602) < 2e-3            # 0.709 * 0.791
    assert abs(1.0 / imp.value_for("+x") - 357.0) < 1.0    # DA: 1/200 -> 1/357
    assert abs(1.0 / imp.value_for("-x") - 357.0) < 1.0
    assert abs(1.0 / imp.value_for("+y") - 535.5) < 1.0    # CA: 1/300 -> 1/535
    # without alpha_hm the EN1993 path stays flat (no reduction, no phi_min lift)
    flat = Imperfection(n_cols=4, phi_s=1/200, standard="EN1993", phi_min=1/500)
    assert abs(flat.value() - 1/200) < 1e-9


def test_member_self_weight_in_dead_case():
    # RSTAB's LC1 has self-weight active (gamma=78.5 kN/m3); OpenSees applies no
    # gravity on its own, so the builder must add A*rho*g member self-weight to
    # the dead case as a global -Z UDL.  Default on; can be switched off.
    from rack15512.builder import _RHO_STEEL, _G_ACC
    m = build_rack(RackConfig(module="single", n_bays=2,
                              levels=[LevelSpec(gap=2000.0)], frame_height=2200.0))
    dead = m.load_cases["dead"]
    # one self-weight UDL per member (plus the dead_load_beam UDLs on the beams)
    sw = [ml for ml in dead.member_loads]
    assert len(sw) >= len(m.members)                  # every member got self-wt
    # an upright's self-weight UDL magnitude = A * rho * g
    up = next(mm for mm in m.members.values() if mm.member_set == "uprights")
    a = m.sections[up.section].A
    got = [abs(ml.qz) for ml in dead.member_loads if ml.member == up.id]
    assert any(abs(q - a * _RHO_STEEL * _G_ACC) < 1e-6 for q in got)
    # off -> no self-weight (only the beam dead UDLs remain)
    m0 = build_rack(RackConfig(module="single", n_bays=2,
                               levels=[LevelSpec(gap=2000.0)], frame_height=2200.0,
                               include_self_weight=False))
    assert len(m0.load_cases["dead"].member_loads) < len(m.members)


def test_calculated_beam_connector_stiffness():
    # 'calculated' beam stiffness = factor * E*I_b/L_b (beam_stiffness module),
    # like the base R899 option; an explicit override still wins.
    from rack15512.beam_stiffness import derived_connector_stiffness

    class _B:                                   # minimal beam section
        Iz = 800000.0
        material = "steel"
    assert abs(derived_connector_stiffness(_B(), 210000.0, 2700.0, 2.0)
               - 2.0 * 210000.0 * 800000.0 / 2700.0) < 1.0
    # factor scales linearly (6 = double-curvature/sway, the stiffest)
    k2 = derived_connector_stiffness(_B(), 210000.0, 2700.0, 2.0)
    k6 = derived_connector_stiffness(_B(), 210000.0, 2700.0, 6.0)
    assert abs(k6 / k2 - 3.0) < 1e-6

    base = dict(module="single", n_bays=2, bay_width=2700.0,
                levels=[LevelSpec(gap=2000.0)], frame_height=2200.0)
    mcalc = build_rack(RackConfig(**base, connector_stiffness_source="calculated",
                                  connector_calc_factor=6.0))
    mman = build_rack(RackConfig(**base, connector_stiffness_source="manual",
                                 connector_stiffness=42.0e6))
    bc = next(mm for mm in mcalc.members.values()
              if mm.member_set == "pallet beams")
    bm = next(mm for mm in mman.members.values()
              if mm.member_set == "pallet beams")
    assert bc.hinge_i.rz > bm.hinge_i.rz          # 6EI/L stiffer than 42 kNm/rad
    assert abs(bm.hinge_i.rz - 42.0e6) < 1.0      # manual value used verbatim


def test_stiffness_gamma_m_softens_design_stiffness():
    # E/gamma_M1 design stiffness (RSTAB "Materials (partial factor gamma_M)"):
    # at 1st order the sway scales as 1/E, so gamma_M=1.1 gives ~1.1x the sway.
    from rack15512.engine.opensees import OpenSeesEngine
    from rack15512.analysis import run_all
    base = dict(module="single", n_bays=2, levels=[LevelSpec(gap=2000.0)],
                frame_height=2200.0)
    m0 = build_rack(RackConfig(**base, stiffness_gamma_m=1.0))
    m1 = build_rack(RackConfig(**base, stiffness_gamma_m=1.1))
    assert m1.analysis.stiffness_gamma_m == 1.1
    for mm in (m0, m1):
        mm.analysis.order = 1                      # linear: clean 1/E scaling

    def sway(model):
        for r in run_all(model):
            if r.converged and r.combo == "ULS1" and r.imp_direction == "+x":
                return max(abs(d[0]) for d in r.displacements.values())
        return 0.0
    s0, s1 = sway(m0), sway(m1)
    assert s0 > 0 and abs(s1 / s0 - 1.1) < 0.02     # ~10% softer


def test_connector_override_applied():
    m = build_rack(RackConfig(module="single", n_bays=2, levels=[LevelSpec(gap=2000.0)],
                              frame_height=2200.0, connector_stiffness_override=73.0e6))
    beams = [mm for mm in m.members.values() if mm.member_set == "pallet beams"]
    assert beams and all(b.hinge_i.rz == 73.0e6 and b.hinge_j.rz == 73.0e6 for b in beams)


def test_base_table_interp_and_tearing():
    tab = [[0, 1e3], [30, 2244e4], [80, 12381e4]]
    f = OpenSeesEngine._interp_base
    assert f(-5, tab) == 1e3                          # uplift -> tearing (C_MIN)
    assert f(100, tab) == 12381e4                     # flat beyond last point
    assert 2244e4 < f(55, tab) < 12381e4              # interpolates


def test_base_table_stored_on_model():
    tab = [[0, 1e3], [60, 8326e4]]
    m = build_rack(RackConfig(module="single", n_bays=2, levels=[LevelSpec(gap=2000.0)],
                              frame_height=2200.0, base_axial_table=tab))
    assert m.base_axial_table == tab


if __name__ == "__main__":
    import pytest
    sys.exit(pytest.main([__file__, "-v"]))


def test_placement_accidental_on_interior_frame_by_default():
    """RSTAB/EN 15512 load the governing INTERIOR upright line (shared between two
    bays), not the corner/edge.  Default load_frame -> an interior line; an
    explicit 0 still pins the end frame; single-bay falls back to the end."""
    m = build_rack(RackConfig(module="single", n_bays=3, levels=[LevelSpec(gap=2000.0)],
                              frame_height=2200.0))
    xs = [nd.x for nd in m.nodes.values()]
    xmin, xmax = min(xs), max(xs)
    pl = m.nodes[m.load_cases["placement"].nodal_loads[0].node]
    ac = m.nodes[m.load_cases["accidental_x"].nodal_loads[0].node]
    assert xmin < pl.x < xmax              # interior, not an end/corner line
    assert xmin < ac.x < xmax

    # explicit end frame still honoured
    e = build_rack(RackConfig(module="single", n_bays=3, levels=[LevelSpec(gap=2000.0)],
                              frame_height=2200.0, load_frame=0))
    assert e.nodes[e.load_cases["placement"].nodal_loads[0].node].x == min(
        nd.x for nd in e.nodes.values())

    # single bay: no interior line -> end frame
    s = build_rack(RackConfig(module="single", n_bays=1, levels=[LevelSpec(gap=2000.0)],
                              frame_height=2200.0))
    assert s.nodes[s.load_cases["placement"].nodal_loads[0].node].x == min(
        nd.x for nd in s.nodes.values())


def test_connector_looseness_modelled_or_lumped():
    # default: looseness recorded on the hinge but carried in phi_l (engine does
    # NOT model the dead-band)
    m0 = build_rack(RackConfig(module="single", n_bays=2, levels=[LevelSpec(gap=2000.0)],
                               frame_height=2200.0, connector_looseness=0.01))
    b0 = next(mm for mm in m0.members.values() if mm.member_set == "pallet beams")
    assert b0.hinge_i.looseness == 0.01                 # recorded
    assert m0.imperfection.phi_l >= 0.01                # carried in the imperfection
    assert m0.model_connector_looseness is False

    # modelled directly: engine models the dead-band, NOT lumped into phi_l
    m1 = build_rack(RackConfig(module="single", n_bays=2, levels=[LevelSpec(gap=2000.0)],
                               frame_height=2200.0, connector_looseness=0.01,
                               model_connector_looseness=True))
    b1 = next(mm for mm in m1.members.values() if mm.member_set == "pallet beams")
    assert b1.hinge_i.looseness == 0.01
    assert m1.imperfection.phi_l == 0.0
    assert m1.model_connector_looseness is True


def test_nonlinear_connector_moment_rotation():
    import os, tempfile
    from rack15512 import io_json
    mphi = [[0.01, 0.73e6], [0.04, 2.0e6], [0.10, 3.0e6]]
    m = build_rack(RackConfig(module="single", n_bays=2, levels=[LevelSpec(gap=2000.0)],
                              frame_height=2200.0, connector_moment_rotation=mphi))
    b = next(mm for mm in m.members.values() if mm.member_set == "pallet beams")
    assert b.hinge_i.m_phi_z == mphi and b.hinge_j.m_phi_z == mphi
    # io_json round-trips the diagram
    p = os.path.join(tempfile.mkdtemp(), "m.json")
    io_json.save(m, p)
    m2 = io_json.load(p)
    b2 = next(mm for mm in m2.members.values() if mm.member_set == "pallet beams")
    assert b2.hinge_i.m_phi_z == mphi


def test_plastic_connector_law():
    import os, tempfile
    from rack15512 import io_json
    m = build_rack(RackConfig(module="single", n_bays=2, levels=[LevelSpec(gap=2000.0)],
                              frame_height=2200.0, connector_plastic=True,
                              connector_m_rd=2.5e6, connector_hardening=0.03,
                              connector_phi_u=0.06))
    b = next(mm for mm in m.members.values() if mm.member_set == "pallet beams")
    assert b.hinge_i.plastic is True
    assert b.hinge_i.m_rd_z == 2.5e6
    assert b.hinge_i.hardening == 0.03 and b.hinge_i.phi_u == 0.06
    # round-trips
    p = os.path.join(tempfile.mkdtemp(), "m.json")
    io_json.save(m, p)
    b2 = next(mm for mm in io_json.load(p).members.values()
              if mm.member_set == "pallet beams")
    assert b2.hinge_i.plastic is True and b2.hinge_i.phi_u == 0.06


def test_rstab_behavior_defaults():
    """RSTAB-matching behaviors (validated on the Zepto model, <=2%):
    SLS combinations linear, EN1993 flat imperfection 1/300 DA / 1/200 CA,
    connector interpolation; the stepped axial-dependent auto base is OPT-IN
    (it can destabilise tall frames and multiplies the solve count)."""
    from rack15512.master_xlsx import load_master
    mw = load_master("examples/Master_Template_FINAL_mount_offset.xlsx")
    kw = dict(master=mw, module="single", n_bays=2,
              bay_width=2300.0, frame_height=2000.0,
              levels=[LevelSpec(gap=1500.0, beam_section="RHS60X40X1.6",
                                pallet_load=5000.0)],
              upright_section="UP0010", steel_E=200000.0, steel_G=76900.0)
    m = build_rack(RackConfig(**kw))
    # SLS combos run geometrically linear (RSTAB), ULS at the model order
    assert all(c.order == 1 for c in m.combinations if c.kind == "SLS")
    assert all(c.order in (None, 2) for c in m.combinations if c.kind == "ULS")
    # default: auto base = single interpolated value (fast, proven); the
    # stepped axial-dependent table only when opted in
    assert m.base_axial_table is None
    m_ax = build_rack(RackConfig(**kw, base_axial_dependent=True))
    assert m_ax.base_axial_table is not None
    assert m_ax.base_axial_table[0][0] == 0.0
    assert m_ax.base_axial_table[1][0] == 30.0
    # EN1993 flat imperfection defaults
    assert m.imperfection.standard == "EN1993"
    assert abs(1 / m.imperfection.value_for("+x") - 300) < 1
    assert abs(1 / m.imperfection.value_for("+y") - 200) < 1
    # steel E/G override
    assert all(mat.E == 200000.0 and mat.G == 76900.0
               for mat in m.materials.values())
    # connector stiffness interpolates between tested upright thicknesses
    s = mw.library.get("RHS60X40X1.6")
    k16, k18, k20 = (s.connector_k_for(t) for t in (1.6, 1.8, 2.0))
    assert k16 < k18 < k20 and abs(k18 - (k16 + k20) / 2) < 1e3
    # order round-trips through json
    import os, tempfile
    from rack15512 import io_json
    p = os.path.join(tempfile.mkdtemp(), "m.json")
    io_json.save(m, p)
    m2 = io_json.load(p)
    assert all(c.order == 1 for c in m2.combinations if c.kind == "SLS")


def test_torsional_restraint_includes_beam_levels():
    """Beam levels restrain twist: L_torsion uses brace nodes AND beam levels,
    while the cross-aisle flexural length L_buckling_y stays on the braces."""
    m = build_rack(RackConfig(module="single", n_bays=2, bay_width=2300.0,
                              levels=[LevelSpec(gap=700.0), LevelSpec(gap=500.0),
                                      LevelSpec(gap=800.0)],
                              frame_height=2200.0, bracing_type="D",
                              bracing_pitch=600.0, bracing_start=150.0))
    ups = [mm for mm in m.members.values() if mm.member_set == "uprights"]
    assert all(mm.L_torsion is not None for mm in ups)
    # twist restraints are a superset of the flexural ones -> never longer
    assert all(mm.L_torsion <= mm.L_buckling_y + 1.0 for mm in ups)
    # somewhere a beam level splits a brace gap -> strictly shorter
    assert any(mm.L_torsion < mm.L_buckling_y - 1.0 for mm in ups)


def test_buckling_interaction_uses_concurrent_station_forces():
    """The buckling interaction takes N, My, Mz from ONE station (concurrent),
    not a mix of maxima; the per-level candidate rows expose the manual
    max-N / max-My / max-Mz method."""
    from rack15512.analysis import run_all
    from rack15512.checks.en15512 import run_checks, upright_set_buckling_rows
    m = build_rack(RackConfig(module="single", n_bays=2, bay_width=2300.0,
                              levels=[LevelSpec(gap=1500.0)],
                              frame_height=1800.0))
    checks = run_checks(m, run_all(m))
    bucks = [c for c in checks if c.check == "BUCKLING" and not c.informative]
    assert bucks
    for c in bucks:
        x = c.extra
        assert "concurrent" in c.detail
        # concurrent values never exceed the member envelopes
        assert x["N"] <= x["N_max"] + 1.0
        assert x["My"] <= x["My_max"] + 1.0
        assert x["Mz"] <= x["Mz_max"] + 1.0
    rows = upright_set_buckling_rows(m, checks)
    assert rows and all(set(r["candidates"]) == {"max N", "max My", "max Mz"}
                        for r in rows)


def test_separate_uls_factors_dl_ll_placement():
    """Each ULS action carries its OWN factor, with the EN 15512 Eq (7)
    multi-variable reduction psi on the SIMULTANEOUS variables: placement
    combos = gamma_G*DL + psi*gamma_Q*LL + psi*gamma_PL*PL.  psi=1.0 applies
    the entered gammas unreduced (flat RSTAB style)."""
    m = build_rack(RackConfig(module="single", n_bays=2,
                              levels=[LevelSpec(gap=2000.0)], frame_height=2200.0,
                              gamma_G=1.2, gamma_Q=1.3, gamma_PL=1.1,
                              multi_var_factor=1.0))
    f1 = next(c.factors for c in m.combinations if c.name == "ULS1")
    f2 = next(c.factors for c in m.combinations if c.name == "ULS2")
    f3 = next(c.factors for c in m.combinations if c.name == "ULS3")
    assert f1 == {"dead": 1.2, "pallets": 1.3}
    assert f2 == {"dead": 1.2, "pallets": 1.3, "placement": 1.1}
    assert f3 == {"dead": 1.2, "pallets": 1.3, "placement_y": 1.1}
    # spec defaults: psi=0.9 and gamma_PL falls back to gamma_Q ->
    # 1.3 DL + 1.26 LL + 1.26 PL (EN 15512 Eq (7): 0.9 x 1.4 = 1.26)
    m2 = build_rack(RackConfig(module="single", n_bays=2,
                               levels=[LevelSpec(gap=2000.0)],
                               frame_height=2200.0))
    f2l = next(c.factors for c in m2.combinations if c.name == "ULS2")
    assert abs(f2l["pallets"] - 1.26) < 1e-9
    assert abs(f2l["placement"] - 1.26) < 1e-9
    assert f2l["dead"] == 1.3
    # direction-locked: placement-X combos analysed with the X imperfection
    c2 = next(c for c in m2.combinations if c.name == "ULS2")
    c3 = next(c for c in m2.combinations if c.name == "ULS3")
    assert c2.imp_directions == ["+x", "-x"]
    assert c3.imp_directions == ["+y", "-y"]
    # SLS carries the sway imperfection, placement excluded (7.3 / R6, R7)
    sls = [c for c in m2.combinations if c.kind == "SLS"]
    assert len(sls) == 1 and sls[0].imp_directions
    assert all("placement" not in k for k in sls[0].factors)


def test_per_role_material_fy_overrides():
    """fy_upright / fy_beam / fy_bracing override the master fy per section
    ROLE; roles sharing a master steel grade are not dragged along; None keeps
    the master value; the role override wins over the fy_override global."""
    from rack15512.master_xlsx import load_master
    mw = load_master("examples/Master_Template_FINAL_mount_offset.xlsx")
    base = dict(master=mw, module="single", n_bays=2, bay_width=2300.0,
                levels=[LevelSpec(gap=1500.0, beam_section="RHS60X40X1.6",
                                  pallet_load=5000.0)],
                frame_height=2000.0, upright_section="UP0010")

    def fy_of(m, name):
        return m.materials[m.sections[name].material].fy

    m = build_rack(RackConfig(**base, fy_upright=250.0, fy_beam=350.0))
    assert fy_of(m, "UP0010") == 250.0            # upright overridden
    assert fy_of(m, "RHS60X40X1.6") == 350.0      # beam overridden
    assert fy_of(m, "1C26X21X1.2") == 270.0       # brace keeps master 270
    m0 = build_rack(RackConfig(**base))           # no overrides -> master fy
    assert fy_of(m0, "UP0010") == 355.0
    assert fy_of(m0, "RHS60X40X1.6") == 270.0
    # role override beats the apply-to-all fy_override global
    m1 = build_rack(RackConfig(**base, steel_fy=355.0, fy_override=True,
                               fy_bracing=235.0))
    assert fy_of(m1, "1C26X21X1.2") == 235.0
    assert fy_of(m1, "RHS60X40X1.6") == 355.0     # global still on the rest


def test_core_checks_only_verdict_scope():
    """core_checks_only: PASS/FAIL from DEFLECTION/STRESS/BUCKLING (+SWAY,
    convergence); connector and other secondary checks become informative but
    are still reported.  Default off keeps every check in the verdict."""
    from rack15512.analysis import run_all
    from rack15512.checks.en15512 import run_checks
    base = dict(module="single", n_bays=2, levels=[LevelSpec(gap=2000.0)],
                frame_height=2200.0)
    m = build_rack(RackConfig(**base, core_checks_only=True))
    checks = run_checks(m, run_all(m))
    for c in checks:
        if c.check in ("STRESS", "BUCKLING", "DEFLECTION", "SWAY"):
            continue                                # core: verdict-carrying
        assert c.informative, c.check               # everything else info-only
    assert any(c.check == "CONNECTOR" for c in checks)   # still reported
    m0 = build_rack(RackConfig(**base))             # default: all checks count
    checks0 = run_checks(m0, run_all(m0))
    assert any(c.check == "CONNECTOR" and not c.informative for c in checks0)


def test_beam_dead_load_from_section_plus_extra_only():
    """The beam self-weight comes from the SELECTED SECTION (A*rho*g); the
    dead_load_beam entry is an ADDITIONAL load only (default 0)."""
    from rack15512.builder import _RHO_STEEL, _G_ACC
    m = build_rack(RackConfig(module="single", n_bays=2,
                              levels=[LevelSpec(gap=2000.0)],
                              frame_height=2200.0))
    bm = next(mm for mm in m.members.values() if mm.member_set == "pallet beams")
    A = m.sections[bm.section].A
    qs = sorted(abs(ml.qz) for ml in m.load_cases["dead"].member_loads
                if ml.member == bm.id)
    assert abs(qs[-1] - A * _RHO_STEEL * _G_ACC) < 1e-6   # self-wt from section
    assert qs[0] < 1e-9                                   # no hardcoded extra
    m2 = build_rack(RackConfig(module="single", n_bays=2,
                               levels=[LevelSpec(gap=2000.0)],
                               frame_height=2200.0, dead_load_beam=0.08))
    bm2 = next(mm for mm in m2.members.values()
               if mm.member_set == "pallet beams")
    qs2 = sorted(abs(ml.qz) for ml in m2.load_cases["dead"].member_loads
                 if ml.member == bm2.id)
    assert any(abs(q - 0.08) < 1e-9 for q in qs2)         # extra entry applied


def test_per_section_fy_override():
    """fy_sections {name: fy}: an entered value overrides the master fy (and
    any per-role override) for THAT section only; unlisted sections keep the
    master / role value."""
    from rack15512.master_xlsx import load_master
    mw = load_master("examples/Master_Template_FINAL_mount_offset.xlsx")
    base = dict(master=mw, module="single", n_bays=2, bay_width=2300.0,
                levels=[LevelSpec(gap=1500.0, beam_section="RHS60X40X1.6",
                                  pallet_load=5000.0)],
                frame_height=2000.0, upright_section="UP0010")

    def fy_of(m, n):
        return m.materials[m.sections[n].material].fy

    m = build_rack(RackConfig(**base, fy_beam=350.0,
                              fy_sections={"RHS60X40X1.6": 330.0,
                                           "UP0010": 280.0}))
    assert fy_of(m, "UP0010") == 280.0            # section beats master 355
    assert fy_of(m, "RHS60X40X1.6") == 330.0      # section beats role 350
    assert fy_of(m, "1C26X21X1.2") == 270.0       # unlisted keeps master
    m0 = build_rack(RackConfig(**base))           # nothing entered -> master
    assert fy_of(m0, "UP0010") == 355.0
    assert fy_of(m0, "RHS60X40X1.6") == 270.0


def test_rstab8_export_tables(tmp_path):
    # The RSTAB 8 export replicates RSTAB's OWN File > Export > Excel
    # structure (validated against an RSTAB 8.29 export of the Zepto
    # model): two header rows, exact sheet names (incl. the trailing
    # space in '1.3 Cross-Sections '), hinge springs in the jy column
    # [kNcm/rad], +/-/value support symbols, wide Factor/No. combination
    # rows with numeric DS codes, per-load-case load sheets and native
    # imperfection LCs.
    import openpyxl
    from rack15512.export_solvers import (to_rstab8_xlsx,
                                          _rstab_section_name, _id_ranges)

    assert _rstab_section_name("RHS100X50X1.6") == "RRO-PAR 100/50/1.6/3.2/1.6/K"
    assert _rstab_section_name("1C36X21X1.2") == "SHAPE-THIN B5H36X21T012"
    assert _rstab_section_name("UP0010") == "SHAPE-THIN UP0010"
    # RSTAB list syntax: dash for runs of 3+, comma for a pair
    assert _id_ranges([3, 1, 2, 7, 9, 10]) == "1-3,7,9,10"
    assert _id_ranges([25, 26]) == "25,26"

    m = build_rack(RackConfig(
        module="single", n_bays=2, bay_width=2300.0, frame_height=5000.0,
        levels=[LevelSpec(gap=1500.0), LevelSpec(gap=1500.0)],
        base_axial_table=[[0, 1.0e3], [30, 3.975e7], [90, 2.5523e8]],
        connector_stiffness=2.039e7))
    path = str(tmp_path / "rstab8.xlsx")
    to_rstab8_xlsx(m, path)
    wb = openpyxl.load_workbook(path)
    for sheet in ("1.1 Nodes", "1.2 Materials", "1.3 Cross-Sections ",
                  "1.4 Member Hinges", "1.7 Members", "1.8 Nodal Supports",
                  "INFO Base Stiffness", "1.11 Sets of Members",
                  "2.1 Load Cases", "2.5 Load Combinations"):
        assert sheet in wb.sheetnames, sheet

    # nodes: RSTAB header pair + Z DOWN (top node at -5000)
    ws = wb["1.1 Nodes"]
    assert ws.cell(2, 1).value == "No." and ws.cell(2, 6).value == "Z [mm]"
    zs = [r[5] for r in ws.iter_rows(min_row=3, values_only=True)]
    assert min(zs) == -5000.0 and max(zs) == 0.0
    # hinge spring sits in the jy column in kNcm/rad (2.039e7 -> 2039)
    hs = list(wb["1.4 Member Hinges"].iter_rows(min_row=3, values_only=True))
    assert any(r[6] == 2039 for r in hs)
    assert all(r[7] == "-" for r in hs)                    # jz released? no - rigid
    # supports: '+' fixed translations, spring value in jY' [kNcm/rad]
    sup = list(wb["1.8 Nodal Supports"].iter_rows(min_row=3,
                                                  values_only=True))[0]
    assert sup[7] == "+" and sup[8] == "+" and sup[9] == "+"
    # jY' carries the LINEAR base spring so it imports AND records directly
    assert isinstance(sup[11], (int, float)) and sup[11] > 0
    assert sup[12] == "+"                # jZ' base-plate torsion FIXED
    # base diagram rows in kN / kNcm/rad incl. the tearing branch
    d = list(wb["INFO Base Stiffness"].iter_rows(min_row=3,
                                              values_only=True))
    assert (30.0, 3975.0) in {(r[3], r[4]) for r in d}
    assert any(r[2] == "PZ'-" for r in d)                  # tearing (uplift)
    # sets: 6 continuous upright lines + per-LEVEL segment sets (one per
    # line per storey: base->L1, L1->L2, L2->top = 3 levels x 6 lines),
    # numbered level-major after the continuous sets (RSTAB reference)
    set_rows = [r for r in wb["1.11 Sets of Members"].iter_rows(
        min_row=3, values_only=True) if r[0]]
    assert len(set_rows) == 6 + 3 * 6
    assert sum(1 for r in set_rows if str(r[1]).startswith("Upright")) == 6
    lvl = [r for r in set_rows if str(r[1]).startswith("LEVEL")]
    assert {r[1] for r in lvl} == {"LEVEL1", "LEVEL2", "LEVEL3"}
    assert min(r[0] for r in lvl) == 7                     # after continuous
    # exactly TWO native imperfection LCs with their own 3.4 sheets
    lcs = list(wb["2.1 Load Cases"].iter_rows(min_row=3, values_only=True))
    imp_lcs = [r for r in lcs if r[3] == "Imperfection"]
    assert len(imp_lcs) == 2
    imp_sheets = [s for s in wb.sheetnames if "3.4 Imperfections" in s]
    # every LC gets a 3.4 sheet (RSTAB export style); only the two
    # imperfection LCs carry rows - alternating +/- per upright set with
    # RSTAB's exact standard string (else the import drops the rows)
    imp_rows = [r for s in imp_sheets
                for r in wb[s].iter_rows(min_row=3, values_only=True) if r[0]]
    assert len(imp_rows) == 4
    assert {r[6] for r in imp_rows} == {300.0, -300.0, 200.0, -200.0}
    assert all(r[4] == "EN 1993-1-1: 2005-07  (Eurocode 3)"
               for r in imp_rows)
    # front upright line of each frame rotated 180 deg; result-combination
    # envelopes per design situation
    rots = {(r[6], r[5]) for r in wb["1.7 Members"].iter_rows(min_row=3,
                                                              values_only=True)
            if r[0]}
    assert (1, 180) in rots and (1, 0) in rots           # both post lines
    # 2.6 Result Combinations is NOT exported: a header-only 2.6 makes
    # RSTAB's import raise a read error, and RSTAB drops the empty table
    assert "2.6 Result Combinations" not in wb.sheetnames
    # empty structural tables exported like RSTAB's own workbook
    for sheet in ("1.5 Member Eccentricities", "1.6 Member Divisions",
                  "1.9 Member Elastic Foundations",
                  "1.10 Member Nonlinearities"):
        assert sheet in wb.sheetnames, sheet
    # combinations: WIDE rows (Factor/No. pairs), numeric DS, one row per
    # (combination x imp direction); the -x twin references the imp LC
    # with factor -1
    co = list(wb["2.5 Load Combinations"].iter_rows(min_row=3,
                                                    values_only=True))
    names = [r[2] for r in co]
    assert "ULS1 (imp +x)" in names and "ULS1 (imp -y)" in names
    assert all(isinstance(r[1], int) for r in co)          # DS codes 1/2/5
    assert any(r[1] == 5 for r in co)                      # SLS present
    row_px = co[names.index("ULS1 (imp +x)")]
    row_mx = co[names.index("ULS1 (imp -x)")]
    pairs_px = list(zip(row_px[4::2], row_px[5::2]))
    pairs_mx = list(zip(row_mx[4::2], row_mx[5::2]))
    assert any(f == 1 and str(lc).startswith("LC") for f, lc in pairs_px)
    assert any(f == -1 for f, lc in pairs_mx)
    # per-load-case load sheets exist (RSTAB export style)
    assert any(s.endswith("3.2 Member Loads") for s in wb.sheetnames)
    # merged group-header cells replicate RSTAB's own export (the import
    # rejects flat headers with 'number of columns is not equal')
    assert {str(r) for r in wb["1.1 Nodes"].merged_cells.ranges} == {"D1:F1"}
    assert {str(r) for r in wb["2.1 Load Cases"].merged_cells.ranges} == \
        {"E1:H1"}
    co_m = {str(r) for r in wb["2.5 Load Combinations"].merged_cells.ranges}
    assert "B1:C1" in co_m and "E1:F1" in co_m
    for s in wb.sheetnames:
        if s.endswith("3.2 Member Loads"):
            assert {str(r) for r in wb[s].merged_cells.ranges} == {"H1:M1"}
    assert "IMPORT NOTES" in wb.sheetnames


def test_staad_deck_complete_and_runnable(tmp_path):
    # The STAAD deck must be self-sufficient: every load case, every
    # combination as a REPEAT LOAD primary case (required for P-Delta) with
    # its sway imperfection as a NOTIONAL LOAD block (factor = gamma * phi
    # per gravity case), one P-Delta analysis command, and rotational
    # springs converted to STAAD's per-DEGREE constants.
    import re as _re
    from rack15512.export_solvers import to_staad

    m = build_rack(RackConfig(
        module="single", n_bays=2, bay_width=2300.0, frame_height=5000.0,
        levels=[LevelSpec(gap=1500.0), LevelSpec(gap=1500.0)],
        connector_stiffness=2.039e7))
    path = str(tmp_path / "deck.std")
    to_staad(m, path)
    txt = open(path).read()

    assert "PDELTA 30 ANALYSIS SMALLDELTA" in txt
    assert "MEMBER TRUSS" in txt                         # braces axial-only
    # members carry the full physical + analysis properties: YD/ZD so
    # STAAD renders the sections, Iy/Iz in APP order (the app's local
    # axes match STAAD's - a horizontal beam's strong axis is IZ)
    pris = [l for l in txt.splitlines() if " PRIS " in l]
    assert pris and all("YD " in l and "ZD " in l for l in pris)
    # shear areas AY/AZ appear when the section provides them (optional)
    import re as _re2
    for l in pris:
        mm = _re2.search(r"IY (\d+\.?\d*) IZ (\d+\.?\d*)", l)
        assert mm
    # per-degree hinge spring: 2.039e7 N*mm/rad * pi/180 = 355872.6 N*mm/deg
    assert _re.search(r"KMZ 355872\.\d", txt)
    # no standalone imperfection load cases - the imperfection lives inside
    # each combination as NOTIONAL LOAD (gamma * phi, signed per direction)
    assert "IMP EHF" not in txt
    ncomb = len(_re.findall(r"^REPEAT LOAD$", txt, _re.M))
    n_named = len(_re.findall(r"^LOAD \d+ LOADTYPE None TITLE", txt, _re.M))
    assert ncomb == n_named and ncomb >= 4
    n_notional = len(_re.findall(r"^NOTIONAL LOAD$", txt, _re.M))
    assert n_notional == ncomb                           # every combo has one
    # ULS1 (imp +x): pallets 1.4 -> notional 1.4/300 = 0.004667 in X;
    # the -x twin flips the sign; cross-aisle uses Z at 1/200
    blk = txt.split("TITLE ULS ULS1 (imp +x)")[1].split("LOAD ")[0]
    assert _re.search(r"X 0\.00466", blk)
    blk = txt.split("TITLE ULS ULS1 (imp -x)")[1].split("LOAD ")[0]
    assert _re.search(r"X -0\.00466", blk)
    blk = txt.split("TITLE ULS ULS1 (imp +y)")[1].split("LOAD ")[0]
    assert _re.search(r"Z 0\.00700", blk)               # 1.4 * 1/200


def test_rstab8_unit_profile_detection_and_conversion(tmp_path):
    # RSTAB interprets imported numbers in the TARGET model's units, so the
    # export adapts: rstab_units_from_export() reads the unit strings from
    # any RSTAB-produced Excel export, and to_rstab8_xlsx(units=...) writes
    # every value and header in that profile.
    import openpyxl
    from rack15512.export_solvers import (to_rstab8_xlsx,
                                          rstab_units_from_export,
                                          _RSTAB_DEFAULT_UNITS)

    m = build_rack(RackConfig(
        module="single", n_bays=2, bay_width=2300.0, frame_height=5000.0,
        levels=[LevelSpec(gap=1500.0), LevelSpec(gap=1500.0)],
        connector_stiffness=2.039e7))
    p0 = str(tmp_path / "default.xlsx")
    to_rstab8_xlsx(m, p0)
    # round trip: our default workbook detects as the factory profile
    assert rstab_units_from_export(p0) == _RSTAB_DEFAULT_UNITS

    prof = {"length": "m", "modulus": "N/mm2", "inertia": "mm4",
            "area": "mm2", "spring_rot": "kNm/rad", "spring_tr": "N/mm",
            "force": "N", "udl": "N/mm", "weight": "t"}
    p1 = str(tmp_path / "custom.xlsx")
    to_rstab8_xlsx(m, p1, units=prof)
    wb = openpyxl.load_workbook(p1)
    # headers carry the profile units and values are converted
    assert wb["1.1 Nodes"].cell(2, 4).value == "X [m]"
    zs = [r[5] for r in wb["1.1 Nodes"].iter_rows(min_row=3,
                                                  values_only=True)]
    assert min(zs) == -5.0                       # 5000 mm -> 5 m
    assert wb["1.2 Materials"].cell(3, 3).value == 210000.0   # E [N/mm2]
    hz = [r[6] for r in wb["1.4 Member Hinges"].iter_rows(min_row=3,
                                                          values_only=True)]
    assert 20.39 in hz                           # 2.039e7 Nmm -> kNm/rad
    ml = next(s for s in wb.sheetnames if s.endswith("3.2 Member Loads"))
    assert wb[ml].cell(2, 8).value == "p [N/mm]"
    # detection round-trips on the custom profile too
    assert rstab_units_from_export(p1) == {**_RSTAB_DEFAULT_UNITS, **prof}


def test_set_member_envelopes_uls():
    # ULS envelope per set of members (RSTAB table 4.2 style): rows for the
    # continuous upright lines AND the per-level storey segments, N most
    # compressive with its CONCURRENT My/Mz, envelopes with governing cases.
    from rack15512.analysis import run_all
    from rack15512.checks.en15512 import set_member_envelopes

    m = build_rack(RackConfig(
        module="single", n_bays=2, bay_width=2300.0, frame_height=5000.0,
        levels=[LevelSpec(gap=1500.0), LevelSpec(gap=1500.0)],
        include_accidental=False, include_pattern=False,
        include_placement=False, mesh_beam=1))
    cases = run_all(m)
    rows = set_member_envelopes(m, cases)
    labels = [r["set"] for r in rows]
    lines = [l for l in labels if " · " not in l]
    segs = [l for l in labels if " · " in l]
    assert len(lines) == 6                      # 3 frames x 2 uprights
    assert len(segs) == 6 * 3                   # base->L1, L1->L2, L2->top
    r0 = next(r for r in rows if " · " not in r["set"])
    assert r0["N_kN"] > 0 and r0["case_N"].startswith("ULS")
    # the line envelope bounds each of its segment envelopes
    for seg in segs:
        line = seg.split(" · ")[0]
        rl = next(r for r in rows if r["set"] == line)
        rs = next(r for r in rows if r["set"] == seg)
        assert rl["N_kN"] >= rs["N_kN"] - 1e-6
        assert rl["Mz_kNm"] >= rs["Mz_kNm"] - 1e-6
